from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base_path = Path(r"D:\Workspace\mini12306\.tmp-report-artifact\template-base.docx")
report_path = Path(r"D:\Workspace\mini12306\实验报告_Mini12306_待提交.docx")
image_dir = Path(r"C:\Users\31794\AppData\Local\Temp")
images = [
    ("codex-clipboard-ed87a57a-73ca-480d-956e-e854f7b8b4c8.png", "图 1  登录页面：手机号与密码登录入口"),
    ("codex-clipboard-69a09ce3-9cfb-42a6-9885-58e476860c4d.png", "图 2  注册页面：虚构实名信息与短信验证码表单"),
    ("codex-clipboard-df40c129-ecbe-4ae7-ae0c-2f8e7c824ac4.png", "图 3  车次列表：车次筛选、排序和余票状态展示"),
    ("codex-clipboard-31742576-6630-493a-b718-60bb68860916.png", "图 4  确认订单：乘车人、席别和费用核对"),
    ("codex-clipboard-090b4cb9-47ea-4a1e-80e9-192d5a5ddffa.png", "图 5  订单支付：模拟支付方式选择与订单金额展示"),
    ("codex-clipboard-4e586da3-22e6-42bd-bb44-ffc5bf05aa8c.png", "图 6  我的订单：不同订单状态及操作入口"),
    ("codex-clipboard-c86608ec-9d2f-40f0-aa87-690b282155bb.png", "图 7  申请退票：退款明细、退票原因与确认操作"),
    ("codex-clipboard-5cda9a0f-f703-40bf-9423-c7feedbdbf26.png", "图 8  车票改签：可选车次、差额计算与确认操作"),
    ("codex-clipboard-f5187146-6edb-4191-9ff5-88069d7c8b3b.png", "图 9  个人中心：旅客专属资料、订单统计和功能入口"),
]
for filename, _ in images:
    if not (image_dir / filename).exists():
        raise FileNotFoundError(image_dir / filename)


def set_font(run, size=10, bold=False, color=None, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def cell_paragraph(cell, text, size=9.5, bold=False, color=None, indent=True, before=0, after=2):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.22)
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return paragraph


def cell_heading(cell, text):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_font(run, size=11, bold=True, color=(26, 91, 196))
    return paragraph


def cell_bullet(cell, text):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run('• ' + text)
    set_font(run, size=8.9)
    return paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    tc_pr.append(shading)


def add_figure_title(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run('Mini-12306（mini01）实验结果：UI 原型截图')
    set_font(run, size=16, bold=True, color=(26, 91, 196), font_name='黑体')
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.space_after = Pt(8)
    intro.paragraph_format.line_spacing = 1.25
    run = intro.add_run('以下截图展示本次实验完成的主要界面。系统使用虚构演示数据，不包含真实身份、银行卡或支付信息。')
    set_font(run, size=10)

# 从干净的模板副本开始，原始文字保持不动，只向两个空白填写区域增加内容。
doc = Document(base_path)
section = doc.sections[-1]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

main_table = doc.tables[0]
process_cell = main_table.rows[4].cells[0]
conclusion_cell = main_table.rows[5].cells[0]

# 第二部分：写入模板中的“实验过程记录”空白区域。
cell_heading(process_cell, '一、系统分析与 UI 原型设计')
cell_paragraph(process_cell, 'Mini-12306（mini01）是一个面向教学与演示的本地铁路购票系统。系统覆盖注册、登录、车次查询、订单确认、模拟支付、订单查看、退票、改签和个人中心；页面采用 375px 移动端优先设计，并可在 Windows、macOS 的主流浏览器中运行。')
cell_paragraph(process_cell, '系统采用“页面—业务—数据”的本地结构：HTML5 负责独立页面，CSS3 负责通用样式和响应式布局，原生 JavaScript 负责输入校验、页面跳转、筛选排序和订单操作，LocalStorage 统一保存虚构用户、车次、库存、订单、查询条件和登录状态。')
cell_bullet(process_cell, '页面层：登录、注册、购票查询、车次列表、订单确认、支付、订单、退票、改签和个人中心页面。')
cell_bullet(process_cell, '数据层：storage.js 统一封装 LocalStorage，mock-data.js 提供 8 条车次和 4 个测试账号。')
cell_bullet(process_cell, '模拟服务：短信、实名认证、支付和退款均为本地模拟流程，不处理真实个人信息与资金。')
cell_heading(process_cell, '二、结对编程过程及任务分工')
cell_paragraph(process_cell, '用户担任驾驶员，负责确认功能范围、页面要求、技术方案和验收结果；Codex 担任领航员，负责根据《Mini-12306 协作规范 v0.3》提供页面结构、样式、JavaScript 逻辑、数据设计、测试思路和问题检查。')
cell_paragraph(process_cell, '本次协作未发生角色互换。驾驶员持续进行需求确认与成果验收；领航员持续进行方案解释、代码实现支持与修改建议。用户确认原生 HTML/CSS/JavaScript + LocalStorage 技术路线，领航员完成项目目录、统一数据访问、模拟数据、登录状态校验、车次查询、订单流程和旅客专属界面的实现支持。')
cell_paragraph(process_cell, '工作照片：待补充。本次未提供真实结对现场照片，报告不虚构照片；提交前可由实验小组补充实际工作照片。', indent=False)

# 第三部分：写入模板中的“结果与讨论”空白区域。
cell_heading(conclusion_cell, '一、开发、测试记录与问题处理')
cell_bullet(conclusion_cell, '完成 css、js、pages 目录结构；公共样式放入 common.css，LocalStorage 读写集中封装在 storage.js。')
cell_bullet(conclusion_cell, '实现手机号与密码登录：手机号为 11 位且以 1 开头；连续输错 3 次锁定 5 分钟；登录成功写入 currentUser；30 分钟无操作要求重新登录。')
cell_bullet(conclusion_cell, '实现查询条件校验与保存、车次列表渲染、全部/高铁动车/普通列车筛选和多种排序；余票以充足、紧张、无票区分，无票车次不可选择。')
cell_bullet(conclusion_cell, '实现订单确认、模拟支付、订单列表、退票、改签和个人中心；订单按 userId 过滤，使不同旅客只查看自己的订单与资料。')
cell_bullet(conclusion_cell, '处理界面问题：移除登录页的底部导航；将订单页退票、改签按钮文字调整为按钮中间对齐。')
cell_bullet(conclusion_cell, '测试结果：测试旅客登录后可查询武汉至北京等线路；有票车次可进入订单确认和模拟支付；订单页可显示不同状态；退票与改签页面可展示金额、规则和确认操作。')
cell_heading(conclusion_cell, '二、实验结论')
cell_paragraph(conclusion_cell, '本实验完成了 Mini-12306 在线购票系统的 UI 原型与核心前端流程演示，满足移动端优先、电脑浏览器可用、页面独立、数据本地持久化及旅客专属展示等目标。系统当前为教学演示版本，数据仅保存在本机浏览器中，实名认证、短信和支付均为模拟流程。后续若要正式运营，需增加服务器、数据库、真实第三方服务和安全合规方案。', size=9.1, after=0)

# 截图作为结果附页，不使用课程要求截图。单独建立新节，确保标题不会落在模板页面底部。
figure_section = doc.add_section(WD_SECTION.NEW_PAGE)
figure_section.top_margin = Inches(0.72)
figure_section.bottom_margin = Inches(0.72)
figure_section.left_margin = Inches(0.78)
figure_section.right_margin = Inches(0.78)
add_figure_title(doc)
for index, (filename, caption) in enumerate(images):
    if index > 0:
        doc.add_page_break()
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(3)
    picture.paragraph_format.space_after = Pt(4)
    picture.add_run().add_picture(str(image_dir / filename), width=Inches(3.18))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(2)
    caption_run = caption_paragraph.add_run(caption)
    set_font(caption_run, size=10.5)

doc.save(report_path)
print(report_path)
